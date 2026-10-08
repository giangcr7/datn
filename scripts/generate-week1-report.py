from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "reports" / "Bao_cao_tuan_1_DiplomaChain.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_bullet(doc, text):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(text, style="List Number")
    p.paragraph_format.space_after = Pt(3)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, value in enumerate(headers):
        hdr[i].text = value
        set_cell_shading(hdr[i], "1F4E78")
        for run in hdr[i].paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.5)
        hdr[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = str(value)
            if row_index % 2 == 1:
                set_cell_shading(cells[i], "EEF4F8")
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    run.font.size = Pt(9.2)
        if len(values) > 0:
            cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            set_cell_border(cell)
            set_cell_margins(cell)
            if widths:
                cell.width = widths[i]
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2.4)
section.right_margin = Cm(2.0)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.15
normal.paragraph_format.space_after = Pt(6)

for style_name, size in (("Title", 20), ("Heading 1", 15), ("Heading 2", 12)):
    style = styles[style_name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(12 if style_name != "Title" else 0)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.keep_with_next = True

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("BÁO CÁO KẾT QUẢ THỰC HIỆN TUẦN 1")
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("Đề tài Hệ thống xác thực văn bằng ứng dụng blockchain")
run.bold = True
run.font.size = Pt(13)

meta = doc.add_table(rows=4, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
meta.autofit = False
metadata = [
    ("Sinh viên thực hiện", "Lê Hoàng Giang"),
    ("Thời gian báo cáo", "Từ ngày 30/09/2026 đến ngày 06/10/2026"),
    ("Tên hệ thống", "DiplomaChain"),
    ("Ngày lập báo cáo", "08/10/2026"),
]
for i, (label, value) in enumerate(metadata):
    meta.cell(i, 0).text = label
    meta.cell(i, 1).text = value
    meta.cell(i, 0).paragraphs[0].runs[0].bold = True
    set_cell_shading(meta.cell(i, 0), "E7EEF4")
    for cell in meta.rows[i].cells:
        set_cell_border(cell)
        set_cell_margins(cell, 100, 120, 100, 120)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(0)
meta.columns[0].width = Cm(4.2)
meta.columns[1].width = Cm(11.2)

doc.add_paragraph()
intro = doc.add_paragraph()
intro.add_run("Kết quả chính  ").bold = True
intro.add_run(
    "Trong tuần 1, em đã kiểm kê mã nguồn, đối chiếu các chức năng với kế hoạch, "
    "chạy kiểm tra build, unit test và lint, lập danh sách lỗi, chốt phạm vi cổng xác thực "
    "doanh nghiệp, xây dựng sơ đồ kiến trúc và chuẩn bị kịch bản demo ban đầu. Mã nguồn và "
    "tài liệu đã được đưa lên GitHub để quản lý phiên bản."
)

doc.add_heading("1 Mục tiêu tuần", level=1)
doc.add_paragraph(
    "Tuần 1 tập trung xác định trạng thái thực tế của hệ thống trước khi phát triển các chức năng mới. "
    "Công việc bao gồm kiểm tra các thành phần Fabric, MongoDB, API và frontend; rà soát các luồng "
    "đăng nhập, cấp văn bằng, xác thực và thu hồi; đồng thời xác định phạm vi cho cổng doanh nghiệp."
)

doc.add_heading("2 Công việc đã thực hiện", level=1)
work_rows = [
    ("1", "Kiểm kê mã nguồn", "Rà soát monorepo gồm NestJS API, Next.js frontend, chaincode, Fabric network và package dùng chung.", "Hoàn thành"),
    ("2", "Rà soát chức năng", "Đối chiếu auth, OTP, cấp, duyệt, từ chối, xác thực, thu hồi, dashboard, explorer và audit log.", "Hoàn thành ở mức mã nguồn"),
    ("3", "Kiểm tra kỹ thuật", "Chạy build, unit test, lint, kiểm tra Docker và các cổng dịch vụ.", "Đã ghi nhận kết quả"),
    ("4", "Phân tích lỗi", "Phân loại lỗi theo mức P0, P1 và P2 để xác định thứ tự xử lý.", "Hoàn thành"),
    ("5", "Chốt cổng doanh nghiệp", "Xác định role verifier, quyền tra cứu, dữ liệu tối thiểu, lịch sử và mã biên nhận.", "Hoàn thành yêu cầu ban đầu"),
    ("6", "Tài liệu hóa", "Tạo báo cáo hiện trạng, sơ đồ kiến trúc, kịch bản demo và smoke-test.", "Hoàn thành"),
    ("7", "Quản lý phiên bản", "Commit và push toàn bộ project lên repository giangcr7/datn, branch main.", "Hoàn thành"),
]
add_table(doc, ["STT", "Công việc", "Nội dung", "Kết quả"], work_rows, [Cm(1.0), Cm(3.2), Cm(8.6), Cm(3.0)])

doc.add_heading("3 Kết quả kiểm tra hệ thống", level=1)
result_rows = [
    ("API unit test", "Đạt", "1 test suite và 1 test đều thành công"),
    ("Shared package", "Đạt", "TypeScript build thành công"),
    ("Chaincode", "Đạt", "TypeScript build thành công"),
    ("NestJS API", "Đạt ở bước build", "Chưa kiểm thử runtime trong lần kiểm kê"),
    ("Next.js frontend", "Chưa đạt", "Build bị chặn do có tiến trình next build khác"),
    ("Lint API", "Chưa đạt", "374 lỗi và 51 cảnh báo, chủ yếu liên quan kiểu any và dữ liệu chưa được định kiểu"),
    ("Docker và Fabric", "Chưa kiểm thử", "Docker Desktop Linux engine chưa hoạt động"),
    ("Smoke-test HTTP", "Chưa đạt", "Frontend và API chưa lắng nghe tại cổng 3000 và 3001"),
]
add_table(doc, ["Hạng mục", "Trạng thái", "Ghi chú"], result_rows, [Cm(4.0), Cm(3.2), Cm(8.6)])

doc.add_heading("4 Chức năng đã có trong mã nguồn", level=1)
for item in [
    "Đăng nhập theo vai trò, JWT, refresh token, đăng xuất một hoặc tất cả thiết bị.",
    "Gửi OTP, kích hoạt tài khoản sinh viên và đổi mật khẩu.",
    "Cấp, yêu cầu, duyệt, từ chối, tìm kiếm và thu hồi văn bằng.",
    "Xác thực văn bằng bằng UUID hoặc dữ liệu proof và đối chiếu hash.",
    "Quản lý sinh viên và nhập dữ liệu từ tệp.",
    "Dashboard thống kê, blockchain explorer, audit log và kiểm tra tính toàn vẹn dữ liệu.",
    "Kết nối Fabric Gateway và chaincode quản lý văn bằng trên Hyperledger Fabric.",
]:
    add_bullet(doc, item)

doc.add_heading("5 Phạm vi cổng xác thực doanh nghiệp", level=1)
doc.add_paragraph(
    "Cổng doanh nghiệp sử dụng role verifier hoặc employer. Tài khoản doanh nghiệp chỉ được tra cứu "
    "và không có quyền sửa dữ liệu nghiệp vụ. Kết quả tra cứu dự kiến gồm thông tin văn bằng tối thiểu, "
    "trạng thái hợp lệ, trạng thái thu hồi và mã biên nhận."
)
for item in [
    "Tra cứu bằng UUID, số hiệu văn bằng hoặc QR.",
    "Ghi lịch sử theo tài khoản, thời gian, phương thức và kết quả.",
    "Giới hạn dữ liệu trả về và áp dụng rate limit theo tài khoản hoặc địa chỉ IP.",
    "Từ chối bằng HTTP 403 nếu doanh nghiệp gọi chức năng cấp, duyệt, từ chối hoặc thu hồi.",
    "Hỗ trợ giao diện desktop và mobile trong giai đoạn triển khai tiếp theo.",
]:
    add_bullet(doc, item)

doc.add_heading("6 Các tồn tại và rủi ro", level=1)
risks = [
    ("P0", "Docker chưa hoạt động", "Chưa thể kiểm thử mạng Fabric và luồng end to end", "Khởi động Docker Desktop và kiểm tra Fabric trước"),
    ("P0", "API và frontend chưa chạy", "Không thể kiểm tra giao diện và endpoint thực tế", "Khởi động dịch vụ, chạy smoke-test"),
    ("P1", "API còn nhiều lỗi lint", "Tăng nguy cơ lỗi runtime và khó bảo trì", "Ưu tiên định kiểu response Fabric và JSON"),
    ("P1", "Test tự động còn ít", "Chưa chứng minh được các luồng nghiệp vụ chính", "Bổ sung unit và integration test"),
    ("P1", "Một số chuỗi tiếng Việt lỗi encoding", "Thông báo có thể hiển thị sai", "Chuẩn hóa toàn bộ mã nguồn về UTF-8"),
    ("P2", "README còn nội dung starter", "Khó cài đặt và bàn giao", "Viết lại hướng dẫn cài đặt, vận hành và demo"),
]
add_table(doc, ["Mức", "Tồn tại", "Ảnh hưởng", "Hướng xử lý"], risks, [Cm(1.4), Cm(4.0), Cm(5.1), Cm(5.3)])

doc.add_heading("7 Thông tin tài khoản phục vụ kiểm thử", level=1)
doc.add_paragraph(
    "Các tài khoản dưới đây được script scripts/start-mongo.js tạo khi cơ sở dữ liệu chưa có người dùng. "
    "Nếu dữ liệu cũ đã tồn tại hoặc mật khẩu đã được đổi, script không ghi đè tài khoản."
)
accounts = [
    ("Quản trị viên", "admin@example.com", "Admin@123", "admin", "Quản trị hệ thống và xem dữ liệu quản trị"),
    ("Nhà trường", "university@example.com", "Univ@123", "university", "Cấp, duyệt, từ chối và thu hồi văn bằng"),
    ("Sinh viên", "student@example.com", "Student@123", "student", "MSSV 2151060001, xem hồ sơ và văn bằng"),
]
add_table(doc, ["Tài khoản", "Email", "Mật khẩu", "Role", "Phạm vi test"], accounts, [Cm(2.5), Cm(4.0), Cm(2.4), Cm(2.2), Cm(4.7)])
doc.add_paragraph(
    "Lưu ý: chức năng đăng nhập kiểm tra đồng thời email, mật khẩu và role. Cần chọn đúng cổng đăng nhập "
    "tương ứng với tài khoản. Hiện chưa có tài khoản doanh nghiệp mẫu trong script seed. Role doanh nghiệp "
    "đang được định nghĩa là employer; phạm vi verifier mới được chốt ở mức yêu cầu."
)

doc.add_heading("8 Hướng dẫn chạy thử", level=1)
for step in [
    "Mở Docker Desktop và chờ Docker engine hoạt động.",
    "Khởi động mạng Hyperledger Fabric theo các script trong thư mục blockchain-network.",
    "Từ thư mục gốc, chạy powershell -ExecutionPolicy Bypass -File .\\start.ps1.",
    "Mở giao diện tại http://localhost:3000 và API Swagger tại http://localhost:3001/docs.",
    "Chạy powershell -ExecutionPolicy Bypass -File .\\scripts\\week1-smoke-test.ps1.",
    "Đăng nhập lần lượt bằng ba tài khoản mẫu và kiểm tra quyền theo vai trò.",
]:
    add_number(doc, step)

doc.add_heading("9 Kế hoạch tuần tiếp theo", level=1)
for item in [
    "Khởi động và kiểm thử đầy đủ môi trường Fabric, MongoDB, API và frontend.",
    "Sửa các lỗi P0, sau đó xử lý nhóm lỗi lint ảnh hưởng trực tiếp đến luồng xác thực.",
    "Bổ sung test cho đăng nhập, cấp, duyệt, xác thực và thu hồi văn bằng.",
    "Chuẩn hóa role doanh nghiệp giữa employer và verifier trước khi triển khai giao diện.",
    "Hoàn thiện README, dữ liệu demo và bằng chứng kiểm thử cho từng chức năng.",
]:
    add_bullet(doc, item)

doc.add_heading("10 Kết luận", level=1)
doc.add_paragraph(
    "Tuần 1 đã hoàn thành mục tiêu khảo sát, kiểm kê và chốt phạm vi phát triển. Hệ thống đã có phần lớn "
    "các module nghiệp vụ chính trong mã nguồn, nhưng cần khôi phục môi trường runtime và xử lý các lỗi "
    "chất lượng trước khi xác nhận hoạt động end to end. Các đầu việc còn lại đã được phân loại để triển "
    "khai theo thứ tự ưu tiên trong tuần tiếp theo."
)

footer = section.footer
p = footer.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Báo cáo tuần 1 DiplomaChain")
r.font.name = "Arial"
r.font.size = Pt(9)
r.font.color.rgb = RGBColor(90, 90, 90)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
